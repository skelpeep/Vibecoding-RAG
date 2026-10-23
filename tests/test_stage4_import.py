"""Meaningful lease/failure/idempotency tests. In-memory external services are explicit."""
from datetime import timedelta
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from atguigu.business import models as m
from atguigu.business.database import make_engine
from atguigu.business.import_errors import ImportFailure, LeaseLost
from atguigu.business.import_parser import parse_document, safe_relative_path, split_document
from atguigu.business.import_service import UploadInput, enqueue_uploads, enqueue_version
from atguigu.business.import_worker import claim_next, complete_task, heartbeat, process_task
from atguigu.config.config import ImportConfig


class MemoryStorage:
    bucket = "test-private"
    def __init__(self):
        self.objects = {}
    def put_bytes(self, key, data, content_type=None, bucket=None):
        self.objects[key] = data
    def get_bytes(self, key, bucket=None):
        return self.objects[key]


class MemoryIndex:
    def __init__(self):
        self.rows = {}
    def ensure_collection(self):
        pass
    def upsert(self, unit_id, version_id, token, chunks, vectors):
        for chunk in chunks:
            self.rows[f"{chunk['id']}:{token}"] = (unit_id, version_id, chunk)
    def verify(self, ids):
        assert set(ids) <= self.rows.keys()


@pytest.fixture
def ctx(tmp_path, monkeypatch):
    engine = make_engine("sqlite:///" + str(tmp_path / "import.db"))
    m.Base.metadata.create_all(engine)
    storage, index = MemoryStorage(), MemoryIndex()
    monkeypatch.setattr("atguigu.business.import_service.get_storage", lambda: storage)
    monkeypatch.setattr(ImportConfig, "retry_base_seconds", 0)
    with Session(engine) as db, db.begin():
        user = m.User(username="importer", display_name="importer", password_hash="test")
        role = m.Role(role_code="importer", role_name="importer")
        db.add_all([user, role]); db.flush()
        for code in ["knowledge:import", "knowledge:create", "knowledge:update"]:
            db.add(m.Permission(code=code, name=code, permission_type="operation")); db.flush()
            db.add(m.RolePermission(role_id=role.id, permission_code=code, permission_type="operation"))
        db.add(m.UserRole(user_id=user.id, role_id=role.id)); db.flush()
        owner_id = user.id
    def enqueue(uploads=None, key=None):
        with Session(engine) as db, db.begin():
            batch = enqueue_uploads(db, owner_id, uploads or [UploadInput("same.md", b"# Heading\nA traceable document.")], key or str(uuid4()))
            return batch.id, list(db.scalars(select(m.ImportTask.id).where(m.ImportTask.batch_id == batch.id)))
    def run(claimed=None, embedding=None):
        claimed = claimed or claim_next(engine)
        assert claimed
        return process_task(*claimed, engine=engine, storage=storage, index=index,
                            embedding=embedding or (lambda texts: {"dense": [[1.0] for t in texts]}), heartbeat_enabled=False)
    yield SimpleNamespace(engine=engine, storage=storage, index=index, owner_id=owner_id, enqueue=enqueue, run=run)
    engine.dispose()


def test_duplicate_names_distinct_ids_and_submission_idempotency(ctx):
    uploads = [UploadInput("same.txt", b"first", relative_path="one/same.txt"), UploadInput("same.txt", b"second", relative_path="two/same.txt")]
    batch_id, tasks = ctx.enqueue(uploads, "same-key")
    repeated_batch, repeated_tasks = ctx.enqueue(uploads, "same-key")
    assert batch_id == repeated_batch and tasks == repeated_tasks
    assert ctx.run() and ctx.run()
    with Session(ctx.engine) as db:
        rows = db.scalars(select(m.ImportTask).where(m.ImportTask.id.in_(tasks))).all()
        assert len({r.unit_id for r in rows}) == len({r.version_id for r in rows}) == 2
        assert db.scalar(select(func.count()).select_from(m.KnowledgeChunk)) == 2
        assert all(db.get(m.KnowledgeVersion, r.version_id).status == "ready" for r in rows)
        assert all(db.get(m.KnowledgeUnit, r.unit_id).active_version_id is None for r in rows)
    from fastapi import HTTPException
    with pytest.raises(HTTPException) as error:
        ctx.enqueue([UploadInput("changed.txt", b"different")], "same-key")
    assert error.value.status_code == 409


def test_mixed_batch_partial_failure(ctx):
    _, tasks = ctx.enqueue([UploadInput("good.txt", b"readable"), UploadInput("legacy.doc", b"old format"), UploadInput("bad.pdf", b"broken")])
    ctx.run(); ctx.run()
    with Session(ctx.engine) as db:
        rows = db.scalars(select(m.ImportTask).where(m.ImportTask.id.in_(tasks))).all()
        assert sorted(r.status for r in rows) == ["completed", "failed", "failed"]
        assert {r.error_code for r in rows} == {None, "LEGACY_DOC_UNSUPPORTED", "INVALID_PDF"}


def test_expired_lease_fences_stale_worker_and_retry_has_stable_chunks(ctx):
    _, (task_id,) = ctx.enqueue()
    old_claim = claim_next(ctx.engine)
    assert claim_next(ctx.engine) is None
    with Session(ctx.engine) as db, db.begin():
        db.get(m.ImportTask, task_id).lease_until = m.utcnow() - timedelta(seconds=1)
    new_claim = claim_next(ctx.engine)
    assert old_claim[1] != new_claim[1]
    with pytest.raises(LeaseLost):
        heartbeat(*old_claim, ctx.engine)
    with pytest.raises(LeaseLost):
        complete_task(*old_claim, "stale", [], [], ctx.engine)
    assert ctx.run(new_claim)
    assert not ctx.run(old_claim)
    with Session(ctx.engine) as db:
        task = db.get(m.ImportTask, task_id)
        assert task.status == "completed" and task.attempts == 2
        assert db.scalar(select(func.count()).select_from(m.KnowledgeChunk)) == 1
        chunk = db.scalar(select(m.KnowledgeChunk))
        assert chunk.vector_id.endswith(new_claim[1])
        assert chunk.id == split_document(parse_document(ctx.storage.get_bytes(db.get(m.StoredFile, task.file_id).object_key), "same.md"), task.version_id)[0]["id"]


def test_retry_exhaustion_keeps_old_active_version_and_no_fake_success(ctx):
    _, (task_id,) = ctx.enqueue()
    assert ctx.run()
    with Session(ctx.engine) as db, db.begin():
        old = db.get(m.ImportTask, task_id)
        unit = db.get(m.KnowledgeUnit, old.unit_id)
        unit.active_version_id, unit.status = old.version_id, "published"
        old_id = old.version_id
        version = m.KnowledgeVersion(unit_id=unit.id, creator_id=ctx.owner_id, version_number=2, content="changed", status="pending")
        db.add(version); db.flush()
        new_task_id = enqueue_version(db, ctx.owner_id, unit, version).id
    def failing(_):
        raise ImportFailure("MODEL_HTTP_503", "真实服务失败", True)
    for _ in range(ImportConfig.max_attempts):
        assert not ctx.run(embedding=failing)
    assert claim_next(ctx.engine) is None
    with Session(ctx.engine) as db:
        task = db.get(m.ImportTask, new_task_id)
        assert task.status == "failed" and task.attempts == task.max_attempts
        assert db.get(m.KnowledgeUnit, task.unit_id).active_version_id == old_id
        assert db.get(m.KnowledgeVersion, old_id).status == "ready"
        assert db.get(m.KnowledgeVersion, task.version_id).status == "failed"


def test_disable_during_embedding_cancels_without_publish(ctx):
    _, (task_id,) = ctx.enqueue()
    def disable(texts):
        with Session(ctx.engine) as db, db.begin():
            task = db.get(m.ImportTask, task_id)
            db.get(m.KnowledgeUnit, task.unit_id).status = "disabled"
        return {"dense": [[1.0] for t in texts]}
    assert not ctx.run(embedding=disable)
    with Session(ctx.engine) as db:
        task = db.get(m.ImportTask, task_id)
        assert task.status == "cancelled"
        assert db.get(m.KnowledgeUnit, task.unit_id).active_version_id is None
        assert db.scalar(select(func.count()).select_from(m.KnowledgeChunk)) == 0


def test_source_integrity_failure(ctx):
    _, (task_id,) = ctx.enqueue()
    with Session(ctx.engine) as db:
        source = db.get(m.StoredFile, db.get(m.ImportTask, task_id).file_id)
        ctx.storage.objects[source.object_key] = b"tampered"
    assert not ctx.run()
    with Session(ctx.engine) as db:
        assert db.get(m.ImportTask, task_id).error_code == "SOURCE_INTEGRITY"


def test_claim_skips_cancelled_and_exhausted_tasks(ctx):
    _, tasks = ctx.enqueue([UploadInput("disabled.txt", b"one"), UploadInput("exhausted.txt", b"two"), UploadInput("valid.txt", b"three")])
    with Session(ctx.engine) as db, db.begin():
        rows = list(db.scalars(select(m.ImportTask).order_by(m.ImportTask.created_at, m.ImportTask.id)))
        db.get(m.KnowledgeUnit, rows[0].unit_id).status = "deleted"
        rows[1].attempts = rows[1].max_attempts
        last_id = rows[2].id
    claimed = claim_next(ctx.engine)
    assert claimed[0] == last_id
    assert ctx.run(claimed)


def test_worker_rechecks_revoked_operation(ctx):
    from sqlalchemy import delete
    _, (task_id,) = ctx.enqueue()
    with Session(ctx.engine) as db, db.begin():
        db.execute(delete(m.RolePermission).where(m.RolePermission.permission_code == "knowledge:import"))
    assert claim_next(ctx.engine) is None
    with Session(ctx.engine) as db:
        task = db.get(m.ImportTask, task_id)
        assert task.status == "cancelled" and task.error_code == "PERMISSION_REVOKED"


def test_pdf_text_and_scan_boundaries(monkeypatch):
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = PdfWriter(); page = writer.add_blank_page(300, 300)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject(); stream.set_data(b"BT /F1 12 Tf 20 200 Td (A PDF paragraph) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO(); writer.write(output)
    parsed = parse_document(output.getvalue(), "text.pdf")
    assert "A PDF paragraph" in parsed.blocks[0].text and parsed.blocks[0].position["page"] == 1
    from PIL import Image
    image = Image.new("RGB", (50, 50), "white"); output = BytesIO(); image.save(output, "PDF")
    monkeypatch.setattr(ImportConfig, "pdf_ocr", False)
    with pytest.raises(ImportFailure) as error:
        parse_document(output.getvalue(), "scan.pdf")
    assert error.value.code == "PDF_OCR_REQUIRED"


def test_private_storage_probe_rejects_anonymous_access():
    from atguigu.business.knowledge_storage import KnowledgeStorage
    from minio.error import S3Error
    storage = object.__new__(KnowledgeStorage)
    storage.bucket = "private"
    class Anonymous:
        def get_object(self, *args, **kwargs):
            return SimpleNamespace(close=lambda: None, release_conn=lambda: None)
    storage.anonymous = Anonymous()
    with pytest.raises(RuntimeError, match="匿名"):
        storage._assert_private("units/test")
    class Denied:
        def get_object(self, *args, **kwargs):
            raise S3Error(response=None, code="AccessDenied", message="denied", resource="", request_id="", host_id="")
    storage.anonymous = Denied()
    storage._assert_private("units/test")


def test_parser_docx_image_and_tables():
    from docx import Document
    from PIL import Image
    picture = BytesIO(); Image.new("RGB", (4, 4), "red").save(picture, "PNG"); picture.seek(0)
    document = Document(); document.add_paragraph("Word paragraph"); document.add_picture(picture)
    table = document.add_table(1, 2); table.cell(0, 0).text = "key"; table.cell(0, 1).text = "value"
    output = BytesIO(); document.save(output)
    parsed = parse_document(output.getvalue(), "word.docx")
    assert len(parsed.images) == 1 and any("key | value" in b.text for b in parsed.blocks)
    assert parsed.images[0].marker in "\n".join(b.text for b in parsed.blocks)


@pytest.mark.parametrize("path", ["../secret.txt", "/etc/passwd", "C:\\private.txt", "x//y.txt", "a/../../b.txt"])
def test_paths_rejected(path):
    with pytest.raises(ImportFailure):
        safe_relative_path(path)


def test_markdown_no_external_image_fetch_and_text_encodings():
    document = parse_document(b"text ![remote](http://127.0.0.1/secret)", "x.md")
    assert "http://" not in document.blocks[0].text and document.warnings
    assert parse_document("中文".encode("gb18030"), "x.txt").blocks[0].text == "中文"


def test_image_inheritance_rewrites_version_file_id(ctx):
    from PIL import Image
    picture = BytesIO(); Image.new("RGB", (4, 4), "red").save(picture, "PNG")
    import base64
    body = b"Text ![embedded](data:image/png;base64," + base64.b64encode(picture.getvalue()) + b")"
    _, (task_id,) = ctx.enqueue([UploadInput("picture.md", body)])
    assert ctx.run()
    with Session(ctx.engine) as db, db.begin():
        task = db.get(m.ImportTask, task_id)
        unit = db.get(m.KnowledgeUnit, task.unit_id)
        old = db.get(m.KnowledgeVersion, task.version_id)
        old_image = db.scalar(select(m.StoredFile).where(m.StoredFile.version_id == old.id, m.StoredFile.purpose == "image"))
        assert old_image.id in old.content
        version = m.KnowledgeVersion(unit_id=unit.id, creator_id=ctx.owner_id, version_number=2, content=old.content, status="pending")
        db.add(version); db.flush()
        enqueue_version(db, ctx.owner_id, unit, version)
        assert old_image.id not in version.content
        new_version_id = version.id
    assert ctx.run()
    with Session(ctx.engine) as db:
        version = db.get(m.KnowledgeVersion, new_version_id)
        assert version.status == "ready" and "/image" in version.content


def test_http_task_metadata_does_not_leak_after_delete_or_to_task_admin(ctx, monkeypatch):
    from fastapi.testclient import TestClient
    from atguigu.business.database import get_db
    from atguigu.business.security import current_user, Principal
    from atguigu.web.api.platform_service import app
    _, (task_id,) = ctx.enqueue()
    def dependency():
        with Session(ctx.engine) as db:
            yield db
            db.commit()
    with Session(ctx.engine) as db:
        user = db.get(m.User, ctx.owner_id)
    principal = Principal(user, set(), {"knowledge:task:read", "knowledge:import"})
    app.dependency_overrides[get_db] = dependency
    app.dependency_overrides[current_user] = lambda: principal
    try:
        with TestClient(app) as client:
            assert client.get(f"/api/knowledge/tasks/{task_id}").json()["file_name"] == "same.md"
            with Session(ctx.engine) as db, db.begin():
                db.get(m.KnowledgeUnit, db.get(m.ImportTask, task_id).unit_id).status = "deleted"
            response = client.get(f"/api/knowledge/tasks/{task_id}").json()
            assert "file_name" not in response and "unit_id" not in response and "error_message" not in response
    finally:
        app.dependency_overrides.clear()


def test_http_multipart_same_names_partial_failure_and_replay(ctx):
    import json
    from fastapi.testclient import TestClient
    from atguigu.business.database import get_db
    from atguigu.business.security import current_user, Principal
    from atguigu.web.api.platform_service import app
    def dependency():
        with Session(ctx.engine) as db:
            try:
                yield db
                db.commit()
            except Exception:
                db.rollback()
                raise
    with Session(ctx.engine) as db:
        user = db.get(m.User, ctx.owner_id)
    principal = Principal(user, set(), {"knowledge:task:read", "knowledge:import"})
    app.dependency_overrides[get_db] = dependency
    app.dependency_overrides[current_user] = lambda: principal
    files = [("files", ("same.txt", b"first", "text/plain")), ("files", ("same.txt", b"second", "text/plain")),
             ("files", ("legacy.doc", b"legacy", "application/msword"))]
    form = {"idempotency_key": "http-key", "relative_paths": json.dumps(["folder-a/same.txt", "folder-b/same.txt", "legacy.doc"])}
    try:
        with TestClient(app) as client:
            response = client.post("/api/knowledge/import", files=files, data=form)
            assert response.status_code == 202, response.text
            batch = response.json()
            assert batch["file_count"] == 3 and batch["counts"]["failed"] == 1 and batch["counts"]["pending"] == 2
            assert len({t["unit_id"] for t in batch["tasks"]}) == 3
            replay = client.post("/api/knowledge/import", files=files, data=form)
            assert replay.status_code == 202 and replay.json()["id"] == batch["id"]
            listing = client.get("/api/knowledge/batches").json()
            assert listing["total"] == 1 and "tasks" not in listing["items"][0]
            assert client.post("/api/knowledge/import", files=files, data={**form, "relative_paths": "[]"}).status_code == 422
            with Session(ctx.engine) as db, db.begin():
                for task in db.scalars(select(m.ImportTask).where(m.ImportTask.batch_id == batch["id"])):
                    db.get(m.KnowledgeUnit, task.unit_id).status = "deleted"
            replay = client.post("/api/knowledge/import", files=files, data=form).json()
            assert all("file_name" not in task and "error_message" not in task for task in replay["tasks"])
    finally:
        app.dependency_overrides.clear()
