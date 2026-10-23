"""Bounded in-memory format adapters; untrusted names never become local paths."""
import base64
import re
import zipfile
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import PurePosixPath
from uuid import NAMESPACE_URL, uuid5

from atguigu.business.import_errors import ImportFailure
from atguigu.config.config import ImportConfig


@dataclass
class ParsedBlock:
    text: str
    position: dict


@dataclass
class ParsedImage:
    name: str
    data: bytes
    content_type: str
    marker: str


@dataclass
class ParsedDocument:
    blocks: list[ParsedBlock]
    images: list[ParsedImage] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def safe_relative_path(value):
    value = (value or "").replace("\\", "/")
    parts = value.split("/")
    if not value or len(value) > 1024 or value.startswith("/") or any(
            p in {"", ".", ".."} or ":" in p or any(ord(c) < 32 for c in p) for p in parts):
        raise ImportFailure("UNSAFE_PATH", "文件相对路径无效，不允许绝对路径、空段或 ..")
    if len(parts[-1]) > 255:
        raise ImportFailure("INVALID_FILENAME", "文件名不能超过 255 个字符")
    return value


def decode_text(data):
    encodings = ["utf-16"] if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ["utf-8-sig", "gb18030"]
    for encoding in encodings:
        try:
            text = data.decode(encoding)
            if "\x00" in text:
                break
            return text.replace("\r\n", "\n").replace("\r", "\n")
        except UnicodeError:
            pass
    raise ImportFailure("TEXT_ENCODING", "文本需使用 UTF-8、带 BOM 的 UTF-16 或 GB18030 编码")


def _bounded(document):
    if len(document.images) > ImportConfig.max_images:
        raise ImportFailure("TOO_MANY_IMAGES", "文档内图片数量超过限制")
    if sum(len(i.data) for i in document.images) > ImportConfig.max_expanded_bytes:
        raise ImportFailure("EXPANDED_SIZE_LIMIT", "文档展开后的图片超出限制")
    if sum(len(b.text) for b in document.blocks) > ImportConfig.max_text_chars:
        raise ImportFailure("TEXT_SIZE_LIMIT", "文档文本长度超过限制")
    if not any(b.text.strip() for b in document.blocks):
        raise ImportFailure("EMPTY_DOCUMENT", "文档未包含可索引正文")
    return document


def _markdown(data):
    text = decode_text(data)
    images, warnings = [], []

    def image(match):
        alt, uri = match.group(1), match.group(2).strip()
        if re.fullmatch(r"/api/knowledge/files/[0-9a-f-]{36}/(?:image|content)", uri):
            # Worker validates each ID against this exact unit/version before indexing.
            return f"![{alt}]({uri})"
        embedded = re.fullmatch(r"data:(image/(?:png|jpeg|gif|webp));base64,([A-Za-z0-9+/=\s]+)", uri)
        if embedded:
            try:
                blob = base64.b64decode(embedded.group(2), validate=True)
            except ValueError:
                raise ImportFailure("INVALID_IMAGE", "Markdown 内嵌图片编码无效")
            marker = f"kb-image:{len(images)}"
            images.append(ParsedImage(f"image-{len(images)}", blob, embedded.group(1), marker))
            return f"![{alt}]({marker})"
        warnings.append("Markdown 外链和相对路径图片未读取；可使用内嵌图片或上传 DOCX/PDF")
        return f"[图片：{alt or '未导入'}]"

    text = re.sub(r"!\[([^\]]*)\]\(([^)]*)\)", image, text)
    # Raw HTML image URLs are never fetched by the parser or exposed as remote images.
    text = re.sub(r"<img\b[^>]*>", "[HTML 图片未导入]", text, flags=re.I)
    return ParsedDocument([ParsedBlock(text, {"format": "markdown", "line_start": 1})], images, sorted(set(warnings)))


def _docx(data):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 10000 or sum(i.file_size for i in members) > ImportConfig.max_expanded_bytes:
                raise ImportFailure("EXPANDED_SIZE_LIMIT", "DOCX 解压大小或文件数量超过限制")
            for member in members:
                safe_relative_path(member.filename.rstrip("/"))
        doc = Document(BytesIO(data))
        blocks, images, seen = [], [], {}
        for index, element in enumerate(doc.iter_inner_content()):
            if isinstance(element, Paragraph):
                text = element.text
                for blip in element._p.xpath(".//a:blip"):
                    rel_id = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                    part = doc.part.related_parts.get(rel_id)
                    if part and part.content_type in {"image/png", "image/jpeg", "image/gif", "image/webp"}:
                        if rel_id not in seen:
                            marker = f"kb-image:{len(images)}"
                            seen[rel_id] = marker
                            images.append(ParsedImage(str(part.partname).rsplit("/", 1)[-1], part.blob, part.content_type, marker))
                        text += f"\n![文档图片]({seen[rel_id]})"
                blocks.append(ParsedBlock(text, {"format": "docx", "paragraph": index + 1}))
            elif isinstance(element, Table):
                blocks.append(ParsedBlock("\n".join(" | ".join(c.text for c in row.cells) for row in element.rows),
                                          {"format": "docx", "table": index + 1}))
        return ParsedDocument(blocks, images)
    except ImportFailure:
        raise
    except Exception as error:
        raise ImportFailure("INVALID_DOCX", "DOCX 文件损坏或格式不受支持") from error


def _pdf(data, name):
    from pypdf import PdfReader
    try:
        reader = PdfReader(BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise ImportFailure("PDF_ENCRYPTED", "加密 PDF 不受支持，请先移除密码")
        if len(reader.pages) > ImportConfig.max_pages:
            raise ImportFailure("PDF_PAGE_LIMIT", "PDF 页数超过限制")
        blocks, images, warnings, scanned = [], [], [], False
        for page_number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            page_images = list(page.images)
            if not text.strip() and page_images:
                scanned = True
            for item in page_images:
                extension = PurePosixPath(item.name).suffix.lower()
                content_type = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}.get(extension)
                if not content_type:
                    warnings.append("部分 PDF 图片编码不受支持，正文保留")
                    continue
                marker = f"kb-image:{len(images)}"
                images.append(ParsedImage(f"page-{page_number}-{item.name}", item.data, content_type, marker))
                text += f"\n![第 {page_number} 页图片]({marker})"
            blocks.append(ParsedBlock(text, {"format": "pdf", "page": page_number}))
            if len(images) > ImportConfig.max_images or sum(len(i.data) for i in images) > ImportConfig.max_expanded_bytes:
                raise ImportFailure("EXPANDED_SIZE_LIMIT", "PDF 图片数量或展开大小超过限制")
            if sum(len(b.text) for b in blocks) > ImportConfig.max_text_chars:
                raise ImportFailure("TEXT_SIZE_LIMIT", "PDF 文本长度超过限制")
        if scanned:
            if ImportConfig.pdf_ocr:
                from atguigu.business.import_mineru import parse_mineru
                return parse_mineru(data, name)
            raise ImportFailure("PDF_OCR_REQUIRED", "PDF 含扫描图片页；当前未启用 OCR，请转为带文本 PDF 或启用 IMPORT_PDF_OCR")
        return ParsedDocument(blocks, images, sorted(set(warnings)))
    except ImportFailure:
        raise
    except Exception as error:
        raise ImportFailure("INVALID_PDF", "PDF 文件损坏或文本提取失败") from error


def parse_document(data, name):
    extension = PurePosixPath(name).suffix.lower()
    if not data:
        raise ImportFailure("EMPTY_FILE", "文件为空")
    if extension in {".md", ".markdown"}:
        result = _markdown(data)
    elif extension == ".txt":
        result = ParsedDocument([ParsedBlock(decode_text(data), {"format": "txt", "line_start": 1})])
    elif extension == ".docx":
        result = _docx(data)
    elif extension == ".pdf":
        result = _pdf(data, name)
    elif extension == ".doc":
        raise ImportFailure("LEGACY_DOC_UNSUPPORTED", "旧版 .doc 暂不支持，请另存为 .docx 后上传")
    else:
        raise ImportFailure("UNSUPPORTED_FORMAT", "支持 PDF、DOCX、Markdown 和 TXT 格式")
    return _bounded(result)


def split_document(document, version_id):
    from langchain_text_splitters import RecursiveCharacterTextSplitter
    splitter = RecursiveCharacterTextSplitter(chunk_size=ImportConfig.chunk_size,
        chunk_overlap=ImportConfig.chunk_overlap, add_start_index=True,
        separators=["\n\n", "\n", "。", "！", "？", "；", ". ", " ", ""])
    chunks = []
    for block_index, block in enumerate(document.blocks):
        for part in splitter.create_documents([block.text]):
            if not part.page_content.strip():
                continue
            start = part.metadata["start_index"]
            chunk_index = len(chunks)
            chunks.append({"id": str(uuid5(NAMESPACE_URL, f"kb:{version_id}:chunk:{chunk_index}")),
                           "chunk_index": chunk_index, "content": part.page_content,
                           "position_json": {**block.position, "block": block_index, "char_start": start,
                                             "char_end": start + len(part.page_content),
                                             "line_start": block.text[:start].count("\n") + 1}})
    if not chunks:
        raise ImportFailure("EMPTY_DOCUMENT", "文档未包含可索引正文")
    return chunks
